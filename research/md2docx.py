import re
from docx import Document
from docx.shared import Pt
src=open('onboarding-plan.md').read().splitlines()
d=Document(); d.styles['Normal'].font.size=Pt(10)
def clean(t):
    t=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',t)
    return t
def addruns(p,t):
    t=clean(t)
    for i,part in enumerate(re.split(r'\*\*',t)):
        r=p.add_run(part.replace('*','') if '`' not in part else part); r.bold=(i%2==1)
i=0;code=False
while i<len(src):
    l=src[i]
    if l.startswith('```'): code=not code; i+=1; continue
    if code: p=d.add_paragraph(); r=p.add_run(l); r.font.name='Courier New'; i+=1; continue
    if l.startswith('|'):
        rows=[]
        while i<len(src) and src[i].startswith('|'):
            if not re.match(r'^\|[-| :]+\|$',src[i]): rows.append([c.strip() for c in src[i].strip('|').split('|')])
            i+=1
        n=max(len(r) for r in rows); t=d.add_table(rows=len(rows),cols=n); t.style='Table Grid'
        for ri,r in enumerate(rows):
            for ci,c in enumerate(r):
                cell=t.cell(ri,ci); cell.text=''; addruns(cell.paragraphs[0],c)
                if ri==0:
                    for run in cell.paragraphs[0].runs: run.bold=True
        continue
    m=re.match(r'^(#+) (.*)',l)
    if m: d.add_heading(clean(m.group(2)),min(len(m.group(1)),4))
    elif l.strip()=='---' or not l.strip(): pass
    elif re.match(r'^\s*[-*] ',l): addruns(d.add_paragraph(style='List Bullet'),re.sub(r'^\s*[-*] ','',l))
    elif re.match(r'^\d+\. ',l): addruns(d.add_paragraph(style='List Number'),re.sub(r'^\d+\. ','',l))
    elif l.startswith('>'): p=d.add_paragraph(); addruns(p,l.lstrip('> ')); p.paragraph_format.left_indent=Pt(18)
    else: addruns(d.add_paragraph(),l)
    i+=1
d.save('onboarding-plan.docx')
